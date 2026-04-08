#!/usr/bin/env python3
"""
map_kernels.py - Map CUDA kernels from nsys profiling data to model layers.

Accepts an nsys-exported SQLite database or CSV file, applies regex-based
mapping rules from a YAML config, and produces an Excel spreadsheet grouped
by model layer.

Usage:
    python map_kernels.py --config mapping_rules.yaml --input profile.sqlite --output results.xlsx
    python map_kernels.py --config mapping_rules.yaml --input profile.csv --output results.xlsx
"""

import argparse
import csv
import os
import re
import sqlite3
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def load_config(config_path: str) -> dict:
    with open(config_path, "r") as f:
        return yaml.safe_load(f)


def detect_input_format(input_path: str) -> str:
    """Auto-detect whether input is SQLite or CSV."""
    with open(input_path, "rb") as f:
        header = f.read(16)
    if header.startswith(b"SQLite format 3"):
        return "sqlite"
    return "csv"


def read_kernels_from_sqlite(db_path: str) -> list[dict]:
    """Read kernel data from nsys SQLite export.

    Returns list of dicts with keys: name, avg_time_us, total_time_us, count.
    """
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    tables = {
        row[0]
        for row in cursor.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        ).fetchall()
    }

    if "CUPTI_ACTIVITY_KIND_KERNEL" in tables:
        query = """
            SELECT
                demangledName AS name,
                AVG(end - start) / 1000.0 AS avg_time_us,
                SUM(end - start) / 1000.0 AS total_time_us,
                COUNT(*) AS count
            FROM CUPTI_ACTIVITY_KIND_KERNEL
            GROUP BY demangledName
            ORDER BY total_time_us DESC
        """
    elif "GPU_KERN_EXEC" in tables:
        # Alternate table name in some nsys versions
        query = """
            SELECT
                demangledName AS name,
                AVG(end - start) / 1000.0 AS avg_time_us,
                SUM(end - start) / 1000.0 AS total_time_us,
                COUNT(*) AS count
            FROM GPU_KERN_EXEC
            GROUP BY demangledName
            ORDER BY total_time_us DESC
        """
    else:
        conn.close()
        print("ERROR: Could not find kernel table in SQLite database.", file=sys.stderr)
        print(f"  Available tables: {sorted(tables)}", file=sys.stderr)
        sys.exit(1)

    rows = cursor.execute(query).fetchall()
    conn.close()

    return [
        {
            "name": row[0],
            "avg_time_us": row[1],
            "total_time_us": row[2],
            "count": row[3],
        }
        for row in rows
    ]


def read_kernels_from_csv(csv_path: str) -> list[dict]:
    """Read kernel data from nsys stats CSV export.

    Handles the `nsys stats --report cuda_gpu_kern_sum --format csv` output,
    which may have leading metadata lines before the actual CSV header.
    """
    with open(csv_path, "r", newline="") as f:
        lines = f.readlines()

    # Skip leading blank/metadata lines until we find the CSV header
    header_idx = None
    for i, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue
        # The header row typically contains "Name" or "Kernel Name"
        if "name" in stripped.lower() and "," in stripped:
            header_idx = i
            break

    if header_idx is None:
        print(
            "ERROR: Could not find CSV header row. Expected a row containing 'Name'.",
            file=sys.stderr,
        )
        sys.exit(1)

    reader = csv.DictReader(lines[header_idx:])
    fieldnames = [f.strip().strip('"') for f in reader.fieldnames]
    reader.fieldnames = fieldnames

    name_col = _find_column(fieldnames, ["Name", "Kernel Name", "name", "kernel_name"])
    avg_col = _find_column(
        fieldnames,
        [
            "Avg (ns)",
            "Avg",
            "avg_ns",
            "Avg Duration",
            "Mean (ns)",
            "Mean",
            "Average",
        ],
    )
    total_col = _find_column(
        fieldnames,
        [
            "Total (ns)",
            "Total",
            "total_ns",
            "Total Duration",
            "Sum (ns)",
            "Sum",
        ],
    )
    count_col = _find_column(
        fieldnames,
        ["Instances", "Count", "instances", "count", "Calls", "calls"],
    )

    if not name_col:
        print(
            f"ERROR: Could not find kernel name column. Available: {fieldnames}",
            file=sys.stderr,
        )
        sys.exit(1)

    kernels = []
    for row in reader:
        name = row.get(name_col, "").strip().strip('"')
        if not name:
            continue

        avg_ns = _parse_number(row.get(avg_col, "0")) if avg_col else 0
        total_ns = _parse_number(row.get(total_col, "0")) if total_col else 0
        count = int(_parse_number(row.get(count_col, "0"))) if count_col else 0

        # nsys CSV reports times in nanoseconds
        kernels.append(
            {
                "name": name,
                "avg_time_us": avg_ns / 1000.0,
                "total_time_us": total_ns / 1000.0,
                "count": count,
            }
        )

    kernels.sort(key=lambda k: k["total_time_us"], reverse=True)
    return kernels


def _find_column(fieldnames: list[str], candidates: list[str]) -> str | None:
    for candidate in candidates:
        for field in fieldnames:
            if field.strip().lower() == candidate.strip().lower():
                return field
    return None


def _parse_number(val: str) -> float:
    val = val.strip().strip('"').replace(",", "")
    if not val or val == "--":
        return 0.0
    return float(val)


def map_kernels_to_layers(
    kernels: list[dict], config: dict
) -> tuple[dict[str, list[dict]], list[dict]]:
    """Apply regex mapping rules to classify kernels into layers.

    Returns:
        mapped: dict mapping layer name -> list of kernel dicts
        unmapped: list of kernel dicts that matched no rule
    """
    layers = config.get("layers", [])

    compiled_rules: list[tuple[dict, list[re.Pattern]]] = []
    for layer in layers:
        patterns = [re.compile(p) for p in layer.get("kernel_patterns", [])]
        compiled_rules.append((layer, patterns))

    mapped: dict[str, list[dict]] = {layer["name"]: [] for layer in layers}
    unmapped: list[dict] = []

    for kernel in kernels:
        matched = False
        for layer, patterns in compiled_rules:
            for pattern in patterns:
                if pattern.search(kernel["name"]):
                    mapped[layer["name"]].append(kernel)
                    matched = True
                    break
            if matched:
                break
        if not matched:
            unmapped.append(kernel)

    return mapped, unmapped


def write_excel(
    mapped: dict[str, list[dict]],
    unmapped: list[dict],
    config: dict,
    output_path: str,
    metadata: dict | None = None,
):
    """Write the mapped kernel data to an Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Kernel Mapping"

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    layer_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Metadata rows
    row = 1
    if metadata:
        for key, val in metadata.items():
            ws.cell(row=row, column=1, value=key).font = Font(bold=True)
            ws.cell(row=row, column=2, value=val)
            row += 1
        row += 1

    model_name = config.get("model", "")
    if model_name:
        ws.cell(row=row, column=1, value=model_name).font = Font(bold=True, size=12)
        row += 1

    # Header row
    headers = ["Model Layer", "Layer Operations", "Kernel", "Time / instance (us)", "Total Time (us)", "Count"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    row += 1

    layers = config.get("layers", [])
    for layer_def in layers:
        layer_name = layer_def["name"]
        operation = layer_def.get("operation", "")
        note = layer_def.get("note", "")
        kernels = mapped.get(layer_name, [])

        if not kernels and not note:
            continue

        if not kernels and note:
            # Layer with no kernels but a note (e.g., fused into another kernel)
            ws.cell(row=row, column=1, value=layer_name).fill = layer_fill
            ws.cell(row=row, column=2, value=operation).fill = layer_fill
            ws.cell(row=row, column=3, value=note)
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row, column=col_idx).border = thin_border
            row += 1
            continue

        first_row = row
        for i, kernel in enumerate(kernels):
            if i == 0:
                ws.cell(row=row, column=1, value=layer_name).fill = layer_fill
                ws.cell(row=row, column=2, value=operation).fill = layer_fill
            ws.cell(row=row, column=3, value=kernel["name"])
            ws.cell(row=row, column=4, value=round(kernel["avg_time_us"], 2)).alignment = Alignment(horizontal="right")
            ws.cell(row=row, column=5, value=round(kernel["total_time_us"], 2)).alignment = Alignment(horizontal="right")
            ws.cell(row=row, column=6, value=kernel["count"]).alignment = Alignment(horizontal="right")
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row, column=col_idx).border = thin_border
            row += 1

        # Merge layer name and operation cells if multiple kernels
        if len(kernels) > 1:
            ws.merge_cells(
                start_row=first_row, start_column=1, end_row=row - 1, end_column=1
            )
            ws.merge_cells(
                start_row=first_row, start_column=2, end_row=row - 1, end_column=2
            )

        # Blank separator row
        row += 1

    # Column widths
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 10

    # Unmapped kernels sheet
    if unmapped:
        ws_unmapped = wb.create_sheet("Unmapped Kernels")
        unmapped_headers = ["Kernel", "Avg Time (us)", "Total Time (us)", "Count"]
        for col_idx, header in enumerate(unmapped_headers, 1):
            cell = ws_unmapped.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = PatternFill(
                start_color="FCE4EC", end_color="FCE4EC", fill_type="solid"
            )
            cell.border = thin_border

        for i, kernel in enumerate(unmapped, start=2):
            ws_unmapped.cell(row=i, column=1, value=kernel["name"])
            ws_unmapped.cell(row=i, column=2, value=round(kernel["avg_time_us"], 2))
            ws_unmapped.cell(row=i, column=3, value=round(kernel["total_time_us"], 2))
            ws_unmapped.cell(row=i, column=4, value=kernel["count"])
            for col_idx in range(1, len(unmapped_headers) + 1):
                ws_unmapped.cell(row=i, column=col_idx).border = thin_border

        ws_unmapped.column_dimensions["A"].width = 80
        ws_unmapped.column_dimensions["B"].width = 18
        ws_unmapped.column_dimensions["C"].width = 18
        ws_unmapped.column_dimensions["D"].width = 10

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(
        description="Map CUDA kernels from nsys profiles to model layers."
    )
    parser.add_argument(
        "--config",
        required=True,
        help="Path to YAML mapping rules config file.",
    )
    parser.add_argument(
        "--input",
        required=True,
        help="Path to nsys-exported SQLite (.sqlite) or CSV file.",
    )
    parser.add_argument(
        "--output",
        default="results.xlsx",
        help="Output Excel file path (default: results.xlsx).",
    )
    parser.add_argument(
        "--label",
        default=None,
        help="Optional label for the run (e.g., 'B200 prefill+decode 1k/8k').",
    )
    args = parser.parse_args()

    if not os.path.exists(args.config):
        print(f"ERROR: Config file not found: {args.config}", file=sys.stderr)
        sys.exit(1)
    if not os.path.exists(args.input):
        print(f"ERROR: Input file not found: {args.input}", file=sys.stderr)
        sys.exit(1)

    config = load_config(args.config)
    fmt = detect_input_format(args.input)

    print(f"Config:  {args.config}")
    print(f"Input:   {args.input} (detected format: {fmt})")
    print(f"Output:  {args.output}")
    print()

    if fmt == "sqlite":
        kernels = read_kernels_from_sqlite(args.input)
    else:
        kernels = read_kernels_from_csv(args.input)

    print(f"Found {len(kernels)} unique kernels.")

    mapped, unmapped = map_kernels_to_layers(kernels, config)

    mapped_count = sum(len(v) for v in mapped.values())
    print(f"Mapped:   {mapped_count} kernels across {sum(1 for v in mapped.values() if v)} layers")
    print(f"Unmapped: {len(unmapped)} kernels")

    if unmapped:
        print("\n--- Unmapped Kernels (review and add patterns to config) ---")
        for k in unmapped[:20]:
            print(f"  {k['avg_time_us']:>10.2f} us  {k['name']}")
        if len(unmapped) > 20:
            print(f"  ... and {len(unmapped) - 20} more (see 'Unmapped Kernels' sheet)")

    metadata = {}
    if args.label:
        metadata["Run"] = args.label
    metadata["Input File"] = Path(args.input).name
    metadata["Model"] = config.get("model", "N/A")

    write_excel(mapped, unmapped, config, args.output, metadata)
    print(f"\nExcel written to: {args.output}")


if __name__ == "__main__":
    main()
